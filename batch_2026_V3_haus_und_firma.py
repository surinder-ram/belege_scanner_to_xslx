import os
import re
import json
import base64
import shutil
import requests
import pandas as pd
import pymupdf  # pip install pymupdf
import datetime

# ===================== CONFIG =====================
API_URL = "http://10.0.0.20:1234/v1/chat/completions"
MODEL_NAME = "qwen3-vl-30b-a3b-instruct"

# 1) QUELL-Verzeichnis (Hier wird nach PDFs gesucht, inkl. aller Unterordner)
BASE_FOLDER = r"H:\.shortcut-targets-by-id\1BqD0qzRaO13EFL7CCxacjy5kR2cXj2Zl\HausverwaltungDaten\03_Hausverwaltung2026\03_Buchhaltung_Steuerberater\2026_Q3"

# 2) ZIEL-Verzeichnis (Hier landen ALLE umbenannten Kopien UND die Excel-Datei)
# Kann ein komplett anderer Pfad sein, z.B. r"D:\Buchhaltung\Export_2026"
TARGET_FOLDER = r"C:\temp\ExportQ3"

# Name der Excel-Datei (landet direkt im TARGET_FOLDER)
OUTPUT_EXCEL_NAME = "_Steuer_Extraktion_Gesamt.xlsx"

# Toleranz in Euro für die Rechenkontrolle (Rundungsdifferenzen)
BETRAG_TOLERANZ = 0.05
MWST_TOLERANZ = 0.10

UMLAUT_MAP = str.maketrans({
    "ä": "ae", "ö": "oe", "ü": "ue", "Ä": "Ae", "Ö": "Oe", "Ü": "Ue", "ß": "ss"
})


# ===================== HILFSFUNKTIONEN =====================

def sanitize_for_filename(text):
    if not text:
        return "Unbekannt"
    text = str(text).translate(UMLAUT_MAP)
    words = re.findall(r"[A-Za-z0-9]+", text)
    return "".join(w[:1].upper() + w[1:] for w in words) or "Unbekannt"


def parse_date_safe(datum_str):
    if not datum_str:
        return None
    match = re.match(r"(\d{1,2})[.\-/](\d{1,2})[.\-/](\d{2,4})", str(datum_str).strip())
    if not match:
        return None
    tag, monat, jahr = match.groups()
    if len(jahr) == 2:
        jahr = "20" + jahr
    try:
        return datetime.date(int(jahr), int(monat), int(tag))
    except ValueError:
        return None


def format_date_for_filename(datum_str):
    d = parse_date_safe(datum_str)
    return d.strftime("%d-%m-%Y") if d else "OhneDatum"


def get_quarter_str(datum_str):
    d = parse_date_safe(datum_str)
    if not d:
        return "Q0"
    quarter = (d.month - 1) // 3 + 1
    return f"Q{quarter}"


def clean_to_float(val):
    if val is None or val == "":
        return 0.0
    s = str(val).replace("EUR", "").replace("€", "").strip()
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "")
        else:
            s = s.replace(",", "")
    s = s.replace(",", ".")
    try:
        match = re.search(r"[-+]?\d*\.\d+|\d+", s)
        return float(match.group(0)) if match else 0.0
    except:
        return 0.0


def extract_embedded_xml(pdf_path):
    try:
        doc = pymupdf.open(pdf_path)
        for i in range(doc.embfile_count()):
            info = doc.embfile_info(i)
            filename = info.get("name", "").lower()
            if "factur-x" in filename or "zugferd" in filename or filename.endswith(".xml"):
                xml_bytes = doc.embfile_get(i)
                return xml_bytes.decode("utf-8", errors="ignore")
    except Exception as e:
        print(f"⚠️ Fehler beim Suchen nach eingebetteten XML-Daten: {e}")
    return None


def extract_pdf_text_if_digital(pdf_path):
    try:
        doc = pymupdf.open(pdf_path)
        full_text = ""
        for page in doc:
            full_text += page.get_text()
        if len(full_text.strip()) > 100:
            return full_text.strip()
    except Exception as e:
        print(f"⚠️ Fehler beim Auslesen des PDF-Texts: {e}")
    return None


def pdf_to_base64_png(pdf_path, target_max_side=1500):
    doc = pymupdf.open(pdf_path)
    page = doc[0]
    rect = page.rect
    orig_max_side = max(rect.width, rect.height)
    zoom = target_max_side / orig_max_side
    zoom = max(1.5, min(zoom, 4.0))
    pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom))
    return base64.b64encode(pix.tobytes("png")).decode("utf-8")


# ===================== VISION-MODELL =====================

EXTRACTION_PROMPT = """Analysiere diesen Beleg (Rechnung/Kassenbon) für die steuerliche Erfassung.
Extrahiere die Daten als JSON. Gib NUR das JSON zurück, keinen weiteren Text.

Hinweise:
- "Land": das Land des Lieferanten/Ausstellers als ISO-Ländercode (z.B. "AT", "DE", "IT").
- "Datum": das Rechnungs-/Belegdatum.
- "Abbuchungsdatum": NUR falls auf dem Beleg selbst ein tatsächliches Zahlungs-/Abbuchungsdatum
  vermerkt ist. Sonst leer ("").
- "MwSt_Satz": Steuersatz in Prozent (z.B. "20%").
- "Kurzbezeichnung": 1-3 Wörter, was gekauft wurde, für einen Dateinamen geeignet.

{
  "Lieferant": "",
  "Land": "",
  "Datum": "TT.MM.JJJJ",
  "Abbuchungsdatum": "TT.MM.JJJJ",
  "Bezeichnung": "",
  "Kurzbezeichnung": "",
  "MwSt_Satz": "",
  "MwSt_Betrag": "",
  "Gesamt": "",
  "Netto": ""
}"""


def build_verification_prompt(vorherige_daten, probleme):
    return f"""Du hast diesen Beleg bereits einmal analysiert. Eine automatische Plausibilitätsprüfung
hat dabei folgende mögliche Probleme gefunden:

{chr(10).join(f"- {p}" for p in probleme)}

Deine vorherige Extraktion war:
{json.dumps(vorherige_daten, ensure_ascii=False, indent=2)}

Schau dir das Bild noch einmal genau an und prüfe gezielt die oben genannten Felder.
Korrigiere falsche Werte. Gib NUR das vollständige JSON zurück.

{{
  "Lieferant": "",
  "Land": "",
  "Datum": "TT.MM.JJJJ",
  "Abbuchungsdatum": "TT.MM.JJJJ",
  "Bezeichnung": "",
  "Kurzbezeichnung": "",
  "MwSt_Satz": "",
  "MwSt_Betrag": "",
  "Gesamt": "",
  "Netto": ""
}}"""


def call_text_model(prompt):
    payload = {
        "model": MODEL_NAME,
        "messages": [{"role": "user", "content": prompt}],
        "temperature": 0.0
    }
    try:
        response = requests.post(API_URL, json=payload, timeout=300)
    except requests.exceptions.RequestException as e:
        print(f"⚠️ Verbindungsfehler (Text-Modell): {e}")
        return None

    if response.status_code != 200:
        return None

    res_data = response.json()
    if "choices" not in res_data:
        return None

    content = res_data["choices"][0]["message"]["content"]
    if not content: return None

    match = re.search(r"\{.*\}", content, re.DOTALL)
    if not match: return None

    json_str = re.sub(r':\s?(\d+)%', r': "\1%"', match.group(0))
    try:
        return json.loads(json_str)
    except:
        return None


def call_vision_model(base64_image, prompt):
    payload = {
        "model": MODEL_NAME,
        "messages": [
            {
                "role": "user",
                "content": [
                    {"type": "text", "text": prompt},
                    {"type": "image_url", "image_url": {"url": f"data:image/png;base64,{base64_image}"}}
                ]
            }
        ],
        "temperature": 0.0
    }
    try:
        response = requests.post(API_URL, json=payload, timeout=600)
    except requests.exceptions.RequestException as e:
        print(f"⚠️ Verbindungsfehler (Vision-Modell): {e}")
        return None

    if response.status_code != 200: return None

    res_data = response.json()
    if "choices" not in res_data: return None

    content = res_data["choices"][0]["message"]["content"]
    if not content: return None

    match = re.search(r"\{.*\}", content, re.DOTALL)
    if not match: return None

    json_str = re.sub(r':\s?(\d+)%', r': "\1%"', match.group(0))
    try:
        return json.loads(json_str)
    except:
        return None


# ===================== QUALITY GATE =====================

def pruefe_plausibilitaet(data):
    issues = []
    if not (data.get("Lieferant") or "").strip(): issues.append("Lieferant fehlt")
    datum = parse_date_safe(data.get("Datum"))
    if datum is None:
        issues.append("Datum fehlt/unlesbar")
    else:
        heute = datetime.date.today()
        if datum.year < 2015 or datum > heute + datetime.timedelta(days=2): issues.append(
            f"Datum unplausibel ({datum})")

    land = (data.get("Land") or "").strip()
    if land and not re.fullmatch(r"[A-Z]{2}", land): issues.append(f"Land-Code unplausibel ({land})")

    brutto = clean_to_float(data.get("Gesamt") or data.get("Gesamtbetrag"))
    netto = clean_to_float(data.get("Netto") or data.get("Nettobetrag"))
    mwst = clean_to_float(data.get("MwSt_Betrag"))

    if brutto == 0.0: issues.append("Bruttobetrag fehlt/0")
    if netto > 0 and mwst > 0 and abs((netto + mwst) - brutto) > BETRAG_TOLERANZ:
        issues.append(f"Netto+MwSt≠Brutto ({netto:.2f}+{mwst:.2f}≠{brutto:.2f})")

    satz_match = re.search(r"[\d.,]+", str(data.get("MwSt_Satz") or ""))
    if satz_match:
        try:
            satz_wert = float(satz_match.group(0).replace(",", "."))
            if not (0 <= satz_wert <= 27):
                issues.append(f"MwSt-Satz unplausibel")
            elif netto > 0 and mwst > 0:
                erwartete_mwst = netto * satz_wert / 100
                if abs(erwartete_mwst - mwst) > MWST_TOLERANZ: issues.append(f"MwSt passt nicht zu Satz")
        except:
            pass

    return issues


def process_beleg(pdf_path):
    filename = os.path.basename(pdf_path)
    print(f"\n{'=' * 50}\n📄 DATEI: {filename}")

    xml_content = extract_embedded_xml(pdf_path)
    data = None

    if xml_content:
        print("⚡ E-Rechnung (XML) gefunden! Nutze XML-Daten...")
        data = call_text_model(EXTRACTION_PROMPT + f"\n\nXML-Rechnungsdaten:\n{xml_content}")

    if not data:
        extracted_text = extract_pdf_text_if_digital(pdf_path)
        if extracted_text:
            print("💡 Digitales PDF erkannt. Nutze direkte Text-Extraktion...")
            data = call_text_model(EXTRACTION_PROMPT + f"\n\nExtrahierter Text:\n{extracted_text}")

    if not data:
        print("📷 Scan erkannt. Nutze Vision-Modell...")
        try:
            base64_image = pdf_to_base64_png(pdf_path)
            data = call_vision_model(base64_image, EXTRACTION_PROMPT)
        except Exception as e:
            print(f"⚠️ Konnte PDF nicht öffnen: {e}")
            return None

    if not data: return None

    issues = pruefe_plausibilitaet(data)
    ki_geprueft = False

    if issues:
        print(f"🔎 Plausibilitätsprüfung meldet Probleme -> frage Modell nochmal...")
        try:
            base64_image = pdf_to_base64_png(pdf_path)
            korrigiert = call_vision_model(base64_image, build_verification_prompt(data, issues))
            if korrigiert:
                data = korrigiert
                ki_geprueft = True
                issues = pruefe_plausibilitaet(data)
        except:
            pass

    status = "OK" if not issues else "PRÜFEN"
    return {"data": data, "status": status, "hinweis": "; ".join(issues), "ki_geprueft": ki_geprueft}


# ===================== EXCEL-EXPORT =====================

def speichere_excel(df, path):
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Belege")
        ws = writer.sheets["Belege"]

        for cell in ws[1]: cell.font = Font(bold=True)
        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "A2"

        for col_idx, col_name in enumerate(df.columns, start=1):
            max_len = max([len(str(col_name))] + [len(str(v)) for v in df[col_name].astype(str)])
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 45)

        status_col_idx = list(df.columns).index("Status") + 1
        gelb = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
        for row_idx in range(2, ws.max_row + 1):
            if ws.cell(row=row_idx, column=status_col_idx).value == "PRÜFEN":
                for cell in ws[row_idx]: cell.fill = gelb


# ===================== MAIN =====================

def main():
    print(f"🚀 Starte Analyse mit {MODEL_NAME} ...")
    alle_ergebnisse = []
    counter = 1

    # Erstelle das definierte Zielverzeichnis
    os.makedirs(TARGET_FOLDER, exist_ok=True)
    target_folder_abs = os.path.abspath(TARGET_FOLDER)

    for root, dirs, files in os.walk(BASE_FOLDER):
        # Schutzschaltung: Falls der TARGET_FOLDER innerhalb des BASE_FOLDER liegt,
        # überspringen wir ihn, damit das Skript nicht in eine Endlosschleife gerät.
        dirs[:] = [d for d in dirs if os.path.abspath(os.path.join(root, d)) != target_folder_abs]

        pdfs = [f for f in files if f.lower().endswith('.pdf')]
        if not pdfs: continue

        ordner_name = os.path.relpath(root, BASE_FOLDER)
        print(f"\n📂 ORDNER: {ordner_name if ordner_name != '.' else 'ROOT'}")

        for f in pdfs:
            pdf_path = os.path.join(root, f)
            ergebnis = process_beleg(pdf_path)
            if not ergebnis:
                print("❌ DATEI ÜBERSPRUNGEN.")
                continue

            data = ergebnis["data"]
            lieferant_clean = sanitize_for_filename(data.get("Lieferant"))
            kurz_clean = sanitize_for_filename(data.get("Kurzbezeichnung") or data.get("Bezeichnung"))
            datum_clean = format_date_for_filename(data.get("Datum"))
            quartal_str = get_quarter_str(data.get("Datum"))

            # Neuer Name
            neuer_name = f"{counter:03d}_{quartal_str}_{lieferant_clean}_{kurz_clean}_{datum_clean}.pdf"

            # Kopiere die Datei in den neuen zentralen TARGET_FOLDER
            ziel_pfad = os.path.join(TARGET_FOLDER, neuer_name)
            shutil.copy2(pdf_path, ziel_pfad)

            alle_ergebnisse.append({
                "Nr": counter,
                "Status": ergebnis["status"],
                "Pruefhinweis": ergebnis["hinweis"],
                "KI_zweifach_geprueft": "Ja" if ergebnis["ki_geprueft"] else "Nein",
                "Quell_Ordner": ordner_name if ordner_name != '.' else 'Hauptordner',
                "Datei_Original": f,
                "Neuer_Dateiname": neuer_name,
                "Quartal": quartal_str,
                "Datum": data.get("Datum"),
                "Abbuchungsdatum": data.get("Abbuchungsdatum"),
                "Lieferant": data.get("Lieferant"),
                "Land": data.get("Land"),
                "Bezeichnung": data.get("Bezeichnung"),
                "Kurzbezeichnung": data.get("Kurzbezeichnung"),
                "MwSt_Satz": data.get("MwSt_Satz"),
                "MwSt_Betrag": clean_to_float(data.get("MwSt_Betrag")),
                "Brutto": clean_to_float(data.get("Gesamt") or data.get("Gesamtbetrag")),
                "Netto": clean_to_float(data.get("Netto") or data.get("Nettobetrag")),
            })

            symbol = "✅" if ergebnis["status"] == "OK" else "⚠️ PRÜFEN"
            print(f"{symbol} | {neuer_name}")
            counter += 1

    if not alle_ergebnisse:
        print("\nKeine Belege gefunden.")
        return

    # Excel-Datei im Root des Ziel-Verzeichnisses speichern
    df = pd.DataFrame(alle_ergebnisse)
    output_path = os.path.join(TARGET_FOLDER, OUTPUT_EXCEL_NAME)
    speichere_excel(df, output_path)

    anzahl_pruefen = int((df["Status"] == "PRÜFEN").sum())
    print(f"\n==================================================")
    print(f"💾 ERFOLG! EXCEL GESPEICHERT IN:")
    print(f"👉 {output_path}")
    print(f"📊 {len(df)} Belege verarbeitet, davon {anzahl_pruefen} zur manuellen Prüfung markiert (gelb).")
    print(f"==================================================")


if __name__ == "__main__":
    main()